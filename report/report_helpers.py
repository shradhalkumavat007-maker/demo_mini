# -*- coding: utf-8 -*-
"""
Report Formatting Helpers for ElectroStock Major Project Report
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(BASE_DIR, "Project Report  CoverPage.docx")
OUTPUT_PATH = os.path.join(BASE_DIR, "Project_Report_ElectroStock.docx")
DIAGRAMS_DIR = os.path.join(BASE_DIR, "diagrams")

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def set_code_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def add_p(doc, text="", font_size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          space_before=0, space_after=4.5, line_spacing=1.16, color_rgb=(20, 20, 20)):
    """
    Adds paragraph(s). Crucially splits on newlines so that soft line-breaks inside
    justified text never trigger wide word spacing. List items are left-aligned.
    """
    if not text:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        return p

    lines = text.split("\n")
    last_p = None
    for idx, raw_line in enumerate(lines):
        line = raw_line.strip()
        if not line:
            continue
        p = doc.add_paragraph()
        is_list = line.startswith(("•", "-", "*")) or (len(line) > 2 and line[0].isdigit() and line[1] in (".", ")"))
        is_short = len(line) < 65

        # Never justify list items or short headings/lines to prevent wide word gaps
        if is_list or is_short:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            p.alignment = align

        if is_list:
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(2.5)
        else:
            p.paragraph_format.space_before = Pt(space_before if idx == 0 else 1.5)
            p.paragraph_format.space_after = Pt(space_after if idx == len(lines) - 1 else 2.5)

        p.paragraph_format.line_spacing = line_spacing
        run = p.add_run(line)
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)
        last_p = p

    return last_p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(15.5)
    run.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_h2(doc, text, space_before=15, space_after=5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3.5)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = True
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT  # Left-aligned: completely prevents word spacing stretching!
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.16

    r_sym = p.add_run("•  ")
    r_sym.font.name = "Times New Roman"
    r_sym.font.size = Pt(12)
    r_sym.bold = True

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.bold = True

    r_txt = p.add_run(text)
    r_txt.font.name = "Times New Roman"
    r_txt.font.size = Pt(12)
    return p

def add_figure(doc, img_filename, caption_text, width=Inches(5.8)):
    img_path = os.path.join(DIAGRAMS_DIR, img_filename)
    if not os.path.exists(img_path):
        print(f"Warning: image {img_path} not found")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(11)
    p_img.paragraph_format.space_after = Pt(3)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=width)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(10.5)
    run_cap.bold = True
    run_cap.font.color.rgb = RGBColor(30, 41, 59)

def add_code_listing(doc, code_lines, caption_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_code_cell_borders(cell)
    set_cell_margins(cell, top=80, bottom=80, left=130, right=130)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05

    for i, line in enumerate(code_lines):
        r = p.add_run(line + ("\n" if i < len(code_lines) - 1 else ""))
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(15, 23, 42)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(11)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)
    r_cap.bold = True
    r_cap.italic = True
    r_cap.font.color.rgb = RGBColor(51, 65, 85)

def add_academic_table(doc, headers, rows, col_widths=None, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(11)
        p_cap.paragraph_format.space_after = Pt(3.5)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.bold = True
        r_cap.font.color.rgb = RGBColor(30, 41, 59)

    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = tbl.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=90, right=90)
        set_cell_borders(hdr_cells[i], top="1E293B", bottom="0F172A", left="334155", right="334155")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=65, bottom=65, left=80, right=80)
            set_cell_borders(row_cells[c_idx], top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0")
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0 or len(str(cell_value)) < 7:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(30, 41, 59)

    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

def setup_headers_and_footers(doc):
    """
    Configures standard academic margins, running header, and perfectly centered
    dynamic PAGE numbers in the footer across all body pages (clean first page).
    """
    for s_idx, section in enumerate(doc.sections):
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True

        # Header on body pages
        header = section.header
        if header.paragraphs:
            p_hdr = header.paragraphs[0]
            p_hdr.text = ""
            p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_hdr.paragraph_format.space_after = Pt(2)
            r_hdr = p_hdr.add_run("ElectroStock: Electronics Inventory & ERP Management System")
            r_hdr.font.name = "Times New Roman"
            r_hdr.font.size = Pt(8.5)
            r_hdr.italic = True
            r_hdr.font.color.rgb = RGBColor(100, 116, 139)

        # Clear existing footer content completely (removes all old lines/borders/right-aligned text)
        footer = section.footer
        footer._element.clear()
        section.first_page_footer._element.clear()

        # Build clean center-aligned footer with dynamic PAGE number field
        p_ftr = footer.add_paragraph()
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ftr.paragraph_format.space_before = Pt(8)
        p_ftr.paragraph_format.space_after = Pt(0)

        fld_xml = (
            f'<w:fldSimple {nsdecls("w")} w:instr="PAGE">'
            f'  <w:r>'
            f'    <w:rPr>'
            f'      <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'      <w:sz w:val="20"/>'
            f'      <w:color w:val="64748B"/>'
            f'    </w:rPr>'
            f'    <w:t>1</w:t>'
            f'  </w:r>'
            f'</w:fldSimple>'
        )
        p_ftr._element.append(parse_xml(fld_xml))
